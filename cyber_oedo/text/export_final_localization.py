"""Export the reviewed final-localization workbook to deterministic dialogue TSV."""

from __future__ import annotations

import argparse
import csv
import hashlib
import pathlib
from collections.abc import Sequence

from cyber_oedo.text.final_dialogue_insertion import (
    DIALOGUE_CATEGORIES,
    FINAL_DIALOGUE_MASTER_SHA256,
    FINAL_LOCALIZATION_WORKBOOK_SHA256,
)

FINAL_SCRIPT_SHEET = "Final Script"
EXPORT_COLUMNS = (
    "ID", "Bank", "Category", "Resource", "Offset", "Japanese Source",
    "Final Localization", "Quality", "Occurrences", "Display Strategy",
)

def export_final_dialogue(workbook: pathlib.Path, output: pathlib.Path) -> str:
    """Write the canonical dialogue-only TSV and return its SHA-256."""
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    if digest != FINAL_LOCALIZATION_WORKBOOK_SHA256:
        raise ValueError(
            "workbook SHA-256 does not match the reviewed 2026-09-24 master"
        )
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError(
            "openpyxl is required; install the editorial extra"
        ) from error
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    if FINAL_SCRIPT_SHEET not in book.sheetnames:
        raise ValueError(f"missing workbook sheet: {FINAL_SCRIPT_SHEET}")
    sheet = book[FINAL_SCRIPT_SHEET]
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator)
    columns = {name: index for index, name in enumerate(header)}
    missing = [name for name in EXPORT_COLUMNS if name not in columns]
    if missing:
        raise ValueError("Final Script is missing columns: " + ", ".join(missing))
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, delimiter="\t", lineterminator="\n")
        writer.writerow(EXPORT_COLUMNS)
        for row in iterator:
            category = row[columns["Category"]]
            if category not in DIALOGUE_CATEGORIES:
                continue
            writer.writerow(
                tuple("" if row[columns[name]] is None else row[columns[name]] for name in EXPORT_COLUMNS)
            )
    result = hashlib.sha256(output.read_bytes()).hexdigest()
    if result != FINAL_DIALOGUE_MASTER_SHA256:
        raise ValueError(
            f"dialogue export hash mismatch: expected {FINAL_DIALOGUE_MASTER_SHA256}, found {result}"
        )
    return result

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Export the reviewed Final Script dialogue rows to TSV."
    )
    parser.add_argument("workbook", type=pathlib.Path)
    parser.add_argument("output", type=pathlib.Path)
    return parser

def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        digest = export_final_dialogue(args.workbook, args.output)
    except (OSError, RuntimeError, ValueError) as error:
        parser.error(str(error))
    print(f"output={args.output.resolve()}")
    print(f"sha256={digest}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
